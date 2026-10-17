"""Build or validate a private manual-annotation workbook; never train from it."""
from __future__ import annotations

import argparse
import json
import math
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from eval.dataset import dataset_fingerprint, file_sha256, load_sources, select_sources  # noqa: E402

DICTIONARY = Path(__file__).with_name("dictionary.json")
TABLES = [ROOT / "research/dataset/inventory.csv", ROOT / "eval/labels.csv", ROOT / "eval/folds.csv"]
VIDEO_SHEET, EPISODE_SHEET = "Видео (32)", "Эпизоды"
VIDEO_HEADERS = [
    "Видео / основной файл", "Все имена-копии / алиасы", "Метка организаторов: 0 / 1",
    "Время организаторов, с", "Тип от организаторов (не указан)", "Длительность, с",
    "SHA256 исходника", "Фолд разработки", "Относительный путь к видео",
    "Тип — человек", "Что именно видно: факты / доказательство", "Начало, с",
    "Конец, с", "Опорное время, с", "Уверенность видимого основания",
    "Комментарий / чего не видно", "Разметчик", "Дата разметки",
]
EPISODE_HEADERS = [
    "Видео / основной файл", "№ записи эпизода", "Тип — человек",
    "Что именно видно: факты / доказательство", "Начало, с", "Конец, с",
    "Опорное время, с", "Уверенность видимого основания", "Комментарий / чего не видно",
    "Разметчик", "Дата разметки",
]


def load_scope():
    all_sources = load_sources(*TABLES)
    sources = select_sources(all_sources, "development")
    if len(sources) != 32 or any(source.fold == 0 for source in sources):
        raise ValueError("Expected exactly 32 development sources and no reserved source")
    if any(row.type_code for source in sources for row in source.rows):
        raise ValueError("Organizer type annotations changed; review this template before rebuilding")
    dictionary = json.loads(DICTIONARY.read_text(encoding="utf-8"))
    return sources, all_sources, dictionary


def source_values(source) -> list:
    return [source.rows[0].video_id, "; ".join(row.video_id for row in source.rows),
            source.has_violation, source.label_secs[0] if len(source.label_secs) == 1 else None,
            None, source.duration_sec, source.sha256, source.fold, source.rows[0].relative_file]


def build(output: Path, episode_rows: int = 256) -> dict:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo

    sources, _, dictionary = load_scope()
    output = output.resolve()
    if output.is_relative_to(ROOT) and not output.is_relative_to(ROOT / "data/annotations"):
        raise ValueError("Private annotation workbooks inside this repository belong in ignored data/annotations/")
    if output.exists():
        raise FileExistsError("Existing annotation workbook will not be overwritten; choose a new output filename")
    if not 32 <= episode_rows <= 1024:
        raise ValueError("Choose 32–1024 empty episode rows")
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Как размечать"
    glossary = workbook.create_sheet("Словарь")
    videos = workbook.create_sheet(VIDEO_SHEET)
    episodes = workbook.create_sheet(EPISODE_SHEET)
    lists = workbook.create_sheet("_Списки")
    codes = [entry["code"] for entry in dictionary["entries"]]
    certainties = [item["code"] for item in dictionary["certainty"]]
    for index, value in enumerate(codes, start=1):
        lists.cell(index, 1, value)
    for index, value in enumerate(certainties, start=1):
        lists.cell(index, 2, value)
    for index, source in enumerate(sources, start=1):
        lists.cell(index, 3, source.rows[0].video_id)
    for index, value in enumerate([code for code in codes if code != "NONE"], start=1):
        lists.cell(index, 4, value)
    for name, column, count in (("AnnotationCodes", "A", len(codes)),
                                ("EvidenceCertainty", "B", len(certainties)),
                                ("DevelopmentSources", "C", len(sources)),
                                ("EpisodeCodes", "D", len(codes) - 1)):
        workbook.defined_names.add(DefinedName(name, attr_text=f"'_Списки'!${column}$1:${column}${count}"))
    lists.sheet_state = "veryHidden"

    instructions.append(["Шаг / поле", "Как заполнять"])
    guidance = [
        ("Назначение", "Ручная разметка видимых действий и оснований для проверки. Это не установление юридического нарушения. Новые записи из книги не используются в обучении автоматически."),
        ("Объём", "Только 32 уникальных источника разработки (33 исходных имени). Восемь отложенных источников не включены: не размечайте и не добавляйте их сейчас."),
        ("1. Сначала просмотр", "Просмотрите запись целиком и контекст до/после подозрительного эпизода. Видео открывает сам разметчик; в этой книге нет кадров, автоматических описаний или предсказаний."),
        ("2. Исходные столбцы", "На листе «Видео (32)» серые столбцы A:I сохранены из таблиц и защищены от случайного редактирования. 0/1 и время организатора — исходные метки; тип организатором не указан. Не заменяйте их своим выводом. SHA, фолд и путь скрыты в G:I, но сохранены."),
        ("3. Один основной тип", "В светлых полях J:R выберите один основной тип из словаря, опишите видимые факты и заполните время/уверенность. Все человеческие поля изначально пусты. Организаторский 0 не подставляет NONE автоматически."),
        ("4. Несколько эпизодов или типов", "На листе «Эпизоды» добавьте отдельную строку для каждого эпизода/типа, выберите ID из списка и дайте уникальный номер записи внутри этого видео. Пересекающиеся интервалы допустимы. Не объединяйте несколько кодов в одной ячейке. Основной эпизод можно повторить здесь: это тот же источник, не независимый пример."),
        ("5. Время", "Все времена — числа секунд от начала исходного видео, допускается дробная часть. Начало ≤ опорное время ≤ конец ≤ длительность. 0 секунд допустимо как настоящее время. Интервал охватывает видимое действие; опорное время выбирается по словарю. Не добавляйте искусственные пороги длительности."),
        ("6. Доказательство", "Коротко опишите, кто что видимо делает, с каким предметом и в какой последовательности. Укажите наблюдаемый контекст, связь с той же суммой/сумкой и доступный обзор. Не угадывайте сумму денег, намерения, должность по одежде или события вне кадра."),
        ("7. Неполный обзор", "Если ключевой этап, предмет, контекст или зона не видны — UNCLEAR и unclear; в комментарии назвать недостающее. Не маскируйте отсутствие обязательного факта словом «возможно» в положительном классе."),
        ("8. NONE", "Только после содержательного полного просмотра, если доступный ракурс позволяет осмысленно оценить действия и оснований из словаря не найдено. Это не юридическая гарантия. Время для NONE оставить пустым; строки эпизода NONE не создавать."),
        ("9. OTHER / UNCLEAR", "Для обоих кодов комментарий обязателен. OTHER — конкретное видимое действие вне словаря; UNCLEAR — не хватает ключевого основания. Подсветка напоминает о пропущенном комментарии; итоговая проверка книги дополнительно проверяет заполнение."),
        ("10. V-OUTSIDER", "Только наблюдение входа в установленную кладовую/сейфовую комнату. Не положительная метка нарушения: полномочия и разрешение по видео неизвестны. Обычный второй человек на стороне кассира не подходит."),
        ("11. Расхождение с организаторами", "Сохраните исходную бинарную метку и время. Свой видимый вывод запишите отдельно; поясните расхождение в комментарии. Не пытайтесь угадать тип по исходному времени без просмотра."),
        ("12. Сохранение", "Сохраните книгу с вашими правками. Генератор никогда не перезаписывает существующую книгу. Не меняйте имена листов, ID и защищённые столбцы; для проверки используется команда из eval/annotation/README.md. Импорта в обучение сейчас нет."),
        ("Уверенность", "Все три значения относятся только к полноте видимого основания, а не к юридической истинности."),
        *[(item["code"], item["label"] + ". " + item["meaning"]) for item in dictionary["certainty"]],
        ("Источник словаря", "rules/sources/S231.yaml; Постановление НБ РК №231. Критерии проверены по источнику, а не скопированы из отключённых автоматических условий. Словарь: eval/annotation/dictionary.json."),
        ("Создано UTC", datetime.now(timezone.utc).isoformat()),
        ("Fingerprint таблиц", dataset_fingerprint(TABLES)),
        ("SHA256 словаря", file_sha256(DICTIONARY)),
    ]
    for row in guidance:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 30
    instructions.column_dimensions["B"].width = 126
    instructions.freeze_panes = "B2"
    for row in range(2, instructions.max_row + 1):
        instructions.row_dimensions[row].height = 60 if row < 16 else 48

    glossary.append(["Код", "Краткий смысл", "Минимальные видимые основания", "Не является достаточным примером",
                     "Интервал и опорное время", "Норма / источник", "Граница вывода", "Вид метки"])
    for entry in dictionary["entries"]:
        glossary.append([entry[key] for key in ("code", "label", "minimum_visible", "non_examples", "timecode", "norm", "boundary", "kind")])
        if entry["code"].startswith("V-"):
            glossary.cell(glossary.max_row, 6).hyperlink = dictionary["source_url"]
    for column, width in zip("ABCDEFGH", [29, 40, 76, 72, 72, 24, 72, 29]):
        glossary.column_dimensions[column].width = width
    for row in range(2, glossary.max_row + 1):
        glossary.row_dimensions[row].height = 150
    glossary.freeze_panes = "B2"
    glossary.auto_filter.ref = glossary.dimensions

    videos.append(VIDEO_HEADERS)
    for source in sources:
        videos.append(source_values(source) + [None] * 9)
        videos.cell(videos.max_row, 1).hyperlink = source.video_path(ROOT / "data/dataset").as_uri()
    episodes.append(EPISODE_HEADERS)
    for _ in range(episode_rows):
        episodes.append([None] * len(EPISODE_HEADERS))

    header_fill, source_fill, editable_fill = (PatternFill("solid", fgColor=color)
                                               for color in ("163B48", "E8EDF1", "F4FAFC"))
    for sheet in (instructions, glossary, videos, episodes):
        sheet.sheet_view.showGridLines = False
        for row in sheet:
            for cell in row:
                cell.font = Font(name="Calibri", size=11, color="172E38")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.fill, cell.font = header_fill, Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        sheet.row_dimensions[1].height = 48
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.print_title_rows = "1:1"
    for row in range(2, 34):
        for column in range(1, 19):
            cell = videos.cell(row, column)
            cell.fill = source_fill if column <= 9 else editable_fill
            cell.protection = Protection(locked=column <= 9)
        videos.row_dimensions[row].height = 80
        for column in (4, 6, 12, 13, 14):
            videos.cell(row, column).number_format = "0.000"
    for row in range(2, episode_rows + 2):
        for cell in episodes[row]:
            cell.protection = Protection(locked=False)
            cell.fill = editable_fill
        episodes.row_dimensions[row].height = 64
        for column in (5, 6, 7):
            episodes.cell(row, column).number_format = "0.000"
    for column, width in zip("ABCDEFGHIJKLMNOPQR", [20, 24, 18, 18, 23, 17, 70, 15, 40, 31, 70, 16, 16, 18, 27, 64, 22, 20]):
        videos.column_dimensions[column].width = width
    for column in "GHI":
        videos.column_dimensions[column].hidden = True
    for column, width in zip("ABCDEFGHIJK", [20, 14, 31, 70, 16, 16, 18, 27, 64, 22, 20]):
        episodes.column_dimensions[column].width = width
    videos.freeze_panes, episodes.freeze_panes = "B2", "C2"
    for sheet, ref, name in ((videos, "A1:R33", "DevelopmentVideoAnnotations"),
                             (episodes, f"A1:K{episode_rows + 1}", "ManualEpisodes")):
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        sheet.add_table(table)
        sheet.auto_filter.ref = ref
        sheet.protection.sheet = True
        sheet.protection.autoFilter = False
        sheet.protection.sort = False
        sheet.protection.selectLockedCells = False
        sheet.protection.selectUnlockedCells = False
    for sheet in (instructions, glossary, lists):
        sheet.protection.sheet = True
        sheet.protection.autoFilter = False

    def list_validation(sheet, cell_range: str, name: str, prompt: str) -> None:
        validation = DataValidation(type="list", formula1=f"={name}", allow_blank=True,
                                    showDropDown=False, showErrorMessage=True, showInputMessage=True)
        validation.errorTitle, validation.error = "Выберите значение из списка", "Свободные коды здесь не допускаются."
        validation.promptTitle, validation.prompt = "Поле ручной разметки", prompt
        validation.errorStyle = "stop"
        sheet.add_data_validation(validation)
        validation.add(cell_range)

    list_validation(videos, "J2:J33", "AnnotationCodes", "Один код. Несколько эпизодов — на отдельном листе.")
    list_validation(videos, "O2:O33", "EvidenceCertainty", "Полнота видимого основания; не юридический вывод.")
    list_validation(episodes, f"A2:A{episode_rows + 1}", "DevelopmentSources", "Только 32 источника разработки.")
    list_validation(episodes, f"C2:C{episode_rows + 1}", "EpisodeCodes", "Одна запись на один эпизод/тип; NONE здесь не используется.")
    list_validation(episodes, f"H2:H{episode_rows + 1}", "EvidenceCertainty", "Полнота видимого основания; не юридический вывод.")
    for sheet, columns, final_row, duration_formula in (
            (videos, "LMN", 33, "$F2"),
            (episodes, "EFG", episode_rows + 1, 'IFERROR(VLOOKUP($A2,\'Видео (32)\'!$A$2:$F$33,6,FALSE),0)')):
        for column in columns:
            formula = f'OR({column}2="",AND(ISNUMBER({column}2),{column}2>=0,{column}2<={duration_formula}))'
            validation = DataValidation(type="custom", formula1=formula, allow_blank=True, showErrorMessage=True)
            validation.errorTitle, validation.errorStyle = "Нужны секунды исходного видео", "stop"
            validation.error = "Введите число от 0 до длительности выбранного источника."
            sheet.add_data_validation(validation)
            validation.add(f"{column}2:{column}{final_row}")
    episode_number = DataValidation(type="whole", operator="between", formula1=1, formula2=9999,
                                    allow_blank=True, showErrorMessage=True, errorStyle="stop")
    episodes.add_data_validation(episode_number)
    episode_number.add(f"B2:B{episode_rows + 1}")
    warning_fill = PatternFill("solid", fgColor="FFF1CC")
    videos.conditional_formatting.add("P2:P33", FormulaRule(
        formula=['AND(OR($J2="OTHER",$J2="UNCLEAR"),LEN(TRIM($P2))=0)'], fill=warning_fill))
    episodes.conditional_formatting.add(f"I2:I{episode_rows + 1}", FormulaRule(
        formula=['AND(OR($C2="OTHER",$C2="UNCLEAR"),LEN(TRIM($I2))=0)'], fill=warning_fill))
    workbook.properties.title = "LADO — ручная разметка 32 источников разработки"
    workbook.properties.subject = "Отдельная человеческая типовая разметка; исходные метки неизменны"
    workbook.properties.creator = "LADO"
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("xb") as stream:
        workbook.save(stream)
    return validate(output, require_empty=True)


def validate(path: Path, *, require_empty: bool = False) -> dict:
    from openpyxl import load_workbook

    sources, all_sources, dictionary = load_scope()
    workbook = load_workbook(path, data_only=False)
    videos, episodes = workbook[VIDEO_SHEET], workbook[EPISODE_SHEET]
    expected = {source.rows[0].video_id: source for source in sources}
    codes = {entry["code"] for entry in dictionary["entries"]}
    certainties = {item["code"] for item in dictionary["certainty"]}
    if videos.max_row != 33 or [cell.value for cell in videos[1]] != VIDEO_HEADERS:
        raise ValueError("The main sheet must retain exactly 32 source rows and its original headers")
    if [cell.value for cell in episodes[1]] != EPISODE_HEADERS:
        raise ValueError("Episode sheet headers changed")
    found, filled_videos, filled_episodes = set(), 0, 0

    def check_human(values: list, duration: float) -> bool:
        if all(value is None for value in values):
            return False
        if require_empty:
            raise ValueError("A new workbook must have no prefilled human annotations")
        code, evidence, start, end, anchor, certainty, comment, _author, _date = values
        if code not in codes or certainty not in certainties or not isinstance(evidence, str) or not evidence.strip():
            raise ValueError("A filled human row needs an allowed type, visible evidence and certainty")
        if code in {"OTHER", "UNCLEAR"} and (not isinstance(comment, str) or not comment.strip()):
            raise ValueError("OTHER and UNCLEAR require a comment")
        if code == "UNCLEAR" and certainty != "unclear":
            raise ValueError("UNCLEAR requires unclear visible certainty")
        times = (start, end, anchor)
        if code == "NONE":
            if any(value is not None for value in times) or certainty != "sufficient_visible":
                raise ValueError("NONE has no episode times and requires a substantive visible review")
        elif code != "UNCLEAR" or any(value is not None for value in times):
            if any(type(value) not in (int, float) or not math.isfinite(value) for value in times) \
                    or not 0 <= start <= anchor <= end <= duration + 1e-6:
                raise ValueError("Episode times must satisfy 0 <= start <= anchor <= end <= duration")
        return True

    for row in videos.iter_rows(min_row=2, max_row=33):
        identity = row[0].value
        if identity not in expected or identity in found:
            raise ValueError("Main source identity changed, repeated or includes an unapproved source")
        found.add(identity)
        source = expected[identity]
        for cell, original in zip(row[:9], source_values(source)):
            equal = math.isclose(cell.value, original, rel_tol=0, abs_tol=1e-9) if type(original) is float and type(cell.value) in (int, float) else cell.value == original
            if not equal or not cell.protection.locked:
                raise ValueError("Organizer source columns changed or lost protection")
        if any(cell.protection.locked for cell in row[9:18]):
            raise ValueError("Human annotation cells must remain editable")
        filled_videos += check_human([cell.value for cell in row[9:18]], source.duration_sec)
    episode_keys = set()
    for row in episodes.iter_rows(min_row=2):
        values = [cell.value for cell in row]
        if all(value is None for value in values):
            continue
        if require_empty:
            raise ValueError("A new episode table must be entirely blank")
        identity, number = values[:2]
        if identity not in expected or type(number) is not int or number < 1 or (identity, number) in episode_keys:
            raise ValueError("Episode row needs an approved source and a unique positive record number")
        if values[2] == "NONE":
            raise ValueError("NONE is not an episode")
        episode_keys.add((identity, number))
        filled_episodes += check_human(values[2:], expected[identity].duration_sec)
    heldout_hashes = {source.sha256 for source in all_sources if source.fold == 0}
    if any(cell.data_type == "f" or (isinstance(cell.value, str) and cell.value in heldout_hashes)
           for sheet in workbook for row in sheet for cell in row):
        raise ValueError("Workbook must contain no cell formulas or reserved source hashes")
    glossary_codes = [workbook["Словарь"].cell(row, 1).value for row in range(2, len(codes) + 2)]
    if glossary_codes != [entry["code"] for entry in dictionary["entries"]]:
        raise ValueError("Dictionary codes changed")
    if not videos.protection.sheet or not episodes.protection.sheet:
        raise ValueError("Worksheet source/header protection was removed")
    workbook.close()
    return {"event": "annotation_workbook_validated", "unique_development_sources": len(found),
            "original_annotation_rows": sum(len(source.rows) for source in sources),
            "reserved_holdout_sources_included": 0, "dictionary_codes": len(codes),
            "filled_human_primary_rows": filled_videos, "filled_episode_rows": filled_episodes,
            "new_human_fields_are_blank": filled_videos == filled_episodes == 0,
            "organizer_type_fields_are_blank": True, "cell_formulas": 0,
            "automatic_training_import": False, "file_sha256": file_sha256(path)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "data/annotations/lado-annotation.xlsx")
    parser.add_argument("--episode-rows", type=int, default=256)
    parser.add_argument("--validate-only", action="store_true")
    parser.add_argument("--require-empty", action="store_true")
    args = parser.parse_args()
    result = validate(args.output, require_empty=args.require_empty) if args.validate_only else build(args.output, args.episode_rows)
    print(json.dumps(result, ensure_ascii=True, sort_keys=True), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
