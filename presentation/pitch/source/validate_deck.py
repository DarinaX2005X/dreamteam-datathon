"""Nonvisual validation for the Lado pitch deck deliverables."""

from __future__ import annotations

import hashlib
import json
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "output" / "lado-pitch.pptx"
PDF = ROOT / "output" / "lado-pitch.pdf"

EXPECTED_LINKS = {
    "https://darinax2005x--lado-ladoservice-web.us-east.modal.run/app",
    "https://github.com/DarinaX2005X/dreamteam-datathon",
}
FORBIDDEN_PATTERNS = {
    "TODO/TBD": re.compile(r"\b(?:TODO|TBD)\b", re.IGNORECASE),
    "overall accuracy claim": re.compile(r"\b(?:accuracy|точност[ьи])\b", re.IGNORECASE),
    "selective score": re.compile(r"(?:2\s*/\s*4|100\s*%)", re.IGNORECASE),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def slide_texts(presentation: Presentation) -> list[str]:
    texts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text)
    return texts


def main() -> int:
    if not PPTX.is_file() or not PDF.is_file():
        raise SystemExit("Missing PPTX or PDF output")

    with zipfile.ZipFile(PPTX) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise SystemExit(f"Corrupt PPTX ZIP member: {bad_member}")
        names = set(archive.namelist())
        note_files = sorted(
            name for name in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)
        )
        slide_files = sorted(
            name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        )
        hyperlinks: set[str] = set()
        for name in names:
            if not re.fullmatch(r"ppt/slides/_rels/slide\d+\.xml\.rels", name):
                continue
            root = ET.fromstring(archive.read(name))
            for relation in root:
                target = relation.attrib.get("Target", "")
                if target.startswith("http"):
                    hyperlinks.add(target)

    presentation = Presentation(PPTX)
    text = "\n".join(slide_texts(presentation))
    forbidden_hits = {
        label: pattern.findall(text)
        for label, pattern in FORBIDDEN_PATTERNS.items()
        if pattern.search(text)
    }

    document = PdfReader(str(PDF))
    pdf_sizes = [
        [round(float(page.mediabox.width), 2), round(float(page.mediabox.height), 2)]
        for page in document.pages
    ]

    result = {
        "pptx_sha256": sha256(PPTX),
        "pdf_sha256": sha256(PDF),
        "pptx_bytes": PPTX.stat().st_size,
        "pdf_bytes": PDF.stat().st_size,
        "slide_count": len(presentation.slides),
        "slide_xml_count": len(slide_files),
        "speaker_note_count": len(note_files),
        "slide_size_inches": [presentation.slide_width / 914400, presentation.slide_height / 914400],
        "pdf_page_count": len(document.pages),
        "pdf_page_sizes_points": pdf_sizes,
        "hyperlinks": sorted(hyperlinks),
        "expected_links_present": EXPECTED_LINKS.issubset(hyperlinks),
        "forbidden_text_hits": forbidden_hits,
        "pptx_zip_integrity": "ok",
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))

    ok = (
        result["slide_count"] == 8
        and result["slide_xml_count"] == 8
        and result["speaker_note_count"] == 8
        and result["pdf_page_count"] == 8
        and result["expected_links_present"]
        and not result["forbidden_text_hits"]
    )
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
